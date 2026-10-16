"""Real OOXML round trips with synthetic data, never a live assessment."""

from copy import deepcopy
from io import BytesIO
import base64
import json
from pathlib import Path
import sys
import unittest
from types import SimpleNamespace
from unittest.mock import Mock
import zipfile
from xml.etree import ElementTree as ET

import pandas as pd
from PIL import Image
from openpyxl import load_workbook
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import resolver_files as f


def scenario():
    datasets = [
        (
            "synthetic.csv",
            pd.DataFrame(
                {
                    "Year": [2001, 2002, 2003, 2004, 2005, 2006],
                    "People": [100, 110, 120, 135, 150, 160],
                    "Urban": [40, 46, 53, 62, 72, 83],
                    "Density": [10, 11, 14, 18, 22, 30],
                }
            ),
        ),
        (
            "synthetic-production.csv",
            pd.DataFrame(
                {
                    "Country": ["A", "B"],
                    "2000/01": [8, 2],
                    "2001/02": [9, 3],
                    "2002/03": [7, 4],
                    "2003/04": [10, 4],
                    "2004/05": [11, 5],
                    "2005/06": [12, 6],
                }
            ),
        ),
    ]
    plan = {
        "version": 1,
        "title": "Synthetic population and supply",
        "steps": ["Review units.", "Compare measured changes."],
        "assumptions": ["Match each season to its ending year."],
        "unresolved": [],
        "sources": [],
        "context": [],
        "images": [],
        "tables": [
            {
                "id": "rates",
                "source": "input1",
                "columns": [
                    {"name": "Year", "expression": "col('Year')"},
                    {"name": "People", "expression": "col('People')"},
                    {
                        "name": "Urban rate",
                        "expression": "coalesce(pct(col('Urban')),at(pct(col('Urban')),1))",
                        "format": "0.00%",
                    },
                    {"name": "Population index", "expression": "index(col('People'))"},
                ],
            },
            {
                "id": "long",
                "operation": "unpivot",
                "source": "input2",
                "id_columns": ["Country"],
                "value_columns": list(datasets[1][1].columns[1:]),
                "variable_name": "Season",
                "value_name": "Output",
            },
            {
                "id": "total",
                "operation": "aggregate",
                "source": "long",
                "group_by": ["Season"],
                "values": [{"column": "Output", "operation": "sum", "name": "Output"}],
            },
            {
                "id": "seasons",
                "source": "total",
                "columns": [
                    {"name": "Year", "expression": "year_end(col('Season'))"},
                    {"name": "Output", "expression": "col('Output')"},
                ],
            },
            {
                "id": "matched",
                "operation": "join",
                "left": "rates",
                "right": "seasons",
                "left_on": "Year",
                "right_on": "Year",
                "how": "inner",
                "validate": "one_to_one",
            },
            {
                "id": "hist",
                "operation": "histogram",
                "source": "input1",
                "column": "Density",
                "bins": 4,
            },
        ],
        "charts": [
            {
                "id": "lines",
                "table": "rates",
                "type": "line",
                "title": "Population and urban growth",
                "x": "Year",
                "x_title": "Year",
                "y_title": "People",
                "y2_title": "Growth",
                "y2_min": 0,
                "y2_max": 1,
                "series": [
                    {"column": "People"},
                    {"column": "Urban rate", "secondary": True},
                ],
            },
            {
                "id": "histogram",
                "table": "hist",
                "type": "column",
                "title": "Density distribution",
                "x": "Median",
                "y_title": "Number of years",
                "gap": 0,
                "data_labels": True,
                "series": [{"column": "Count"}],
            },
            {
                "id": "scatter",
                "table": "matched",
                "type": "scatter",
                "title": "Population and supply",
                "x": "People",
                "x_title": "People",
                "y_title": "Output (source units)",
                "series": [{"column": "Output"}],
            },
            {
                "id": "area",
                "table": "rates",
                "type": "area",
                "title": "Population index",
                "x": "Year",
                "series": [{"column": "Population index"}],
            },
            {
                "id": "bars",
                "table": "input2",
                "type": "bar",
                "title": "Production by country",
                "x": "Country",
                "series": [{"column": "2000/01"}],
            },
            {
                "id": "pie",
                "table": "input2",
                "type": "pie",
                "title": "Production share",
                "x": "Country",
                "series": [{"column": "2000/01"}],
            },
        ],
    }
    return datasets, plan


class FileWorkflowTests(unittest.TestCase):
    def test_export_check_rejects_corrupted_chart_values(self):
        data, plan = scenario()
        tables, _ = f.execute_plan(plan, data)
        blob = f.build_excel(plan, tables)
        output = BytesIO()
        with (
            zipfile.ZipFile(BytesIO(blob)) as source,
            zipfile.ZipFile(output, "w") as target,
        ):
            for item in source.infolist():
                value = source.read(item.filename)
                if item.filename == "xl/charts/chart1.xml":
                    chart = ET.fromstring(value)
                    chart.find(
                        f".//{{{f.C}}}val/{{{f.C}}}numRef/{{{f.C}}}numCache/{{{f.C}}}pt/{{{f.C}}}v"
                    ).text = "999999"
                    value = ET.tostring(chart)
                target.writestr(item, value)
        with self.assertRaisesRegex(f.PlanError, "chart cache differs"):
            f.validate_files(plan, tables, output.getvalue())

    def test_sort_filter_and_blank_formula_policy(self):
        data = [
            (
                "unsorted.csv",
                pd.DataFrame({"Year": [2022, 2020, 2021], "Value": [30, 10, 20]}),
            )
        ]
        plan = f.starter_plan(data)
        plan["tables"] = [
            {
                "id": "sorted",
                "source": "input1",
                "sort_by": ["Year"],
                "columns": [
                    {"name": "Year", "expression": "col('Year')"},
                    {
                        "name": "Growth",
                        "expression": "pct(col('Value'))",
                        "format": "0.00%",
                    },
                ],
            },
            {
                "id": "filtered",
                "source": "sorted",
                "where": "col('Year') >= 2021",
                "columns": [
                    {"name": "Year", "expression": "col('Year')"},
                    {"name": "Growth", "expression": "col('Growth')"},
                ],
            },
        ]
        plan["charts"][0].update(table="filtered", series=[{"column": "Growth"}])
        tables, _ = f.execute_plan(plan, data)
        self.assertEqual(tables["filtered"].frame.Growth.tolist(), [1, 0.5])
        self.assertTrue(pd.isna(tables["sorted"].frame.Growth.iloc[0]))
        f.validate_files(plan, tables, f.build_excel(plan, tables))

    def test_complete_calculation_and_office_round_trip(self):
        data, plan = scenario()
        tables, checks = f.execute_plan(plan, data)
        self.assertEqual(
            tables["total"].frame.Output.tolist(), [10, 12, 11, 14, 16, 18]
        )
        self.assertEqual(
            tables["matched"].frame.Year.tolist(), [2001, 2002, 2003, 2004, 2005, 2006]
        )
        self.assertEqual(tables["hist"].frame.Count.sum(), 6)
        self.assertAlmostEqual(tables["rates"].frame["Urban rate"][0], 0.15)
        workbook = f.build_excel(plan, tables)
        deck = f.build_powerpoint(plan, tables, workbook, {})
        checks += f.validate_files(plan, tables, workbook, deck)
        self.assertTrue(all(c["status"] in ("pass", "manual") for c in checks))
        self.assertEqual(len(Presentation(BytesIO(deck)).slides), 12)
        with zipfile.ZipFile(BytesIO(deck)) as package:
            chart = ET.fromstring(package.read("ppt/charts/chart1.xml"))
            self.assertEqual(len(chart.findall(f".//{{{f.C}}}valAx")), 2)
            self.assertEqual(len(chart.findall(f".//{{{f.C}}}ser")), 2)
            self.assertIn(b"Urban rate", package.read("ppt/charts/chart1.xml"))
            embedded = [
                p for p in package.namelist() if p.startswith("ppt/embeddings/")
            ]
            self.assertEqual(len(embedded), 6)
            for p in embedded:
                self.assertEqual(package.read(p), workbook)
        guide = f.instructions(plan, tables, checks)
        previews = {
            c["id"]: f.chart_preview(c, tables[c["table"]]) for c in plan["charts"]
        }
        for blob in previews.values():
            with Image.open(BytesIO(blob)) as im:
                self.assertGreater(im.width, 1000)
        pack = f.review_pack(workbook, deck, previews, {}, guide, checks)
        with zipfile.ZipFile(BytesIO(pack)) as archive:
            self.assertIn("Presentation.pptx", archive.namelist())
            self.assertIn("Step-by-step guide.md", archive.namelist())
            self.assertFalse(any("prompt" in n.lower() for n in archive.namelist()))

    def test_full_rows_not_preview_and_source_injection_is_literal(self):
        data = [
            (
                "unsafe.csv",
                pd.DataFrame(
                    {
                        "Label": ['=HYPERLINK("https://example.test")']
                        + list(map(str, range(1, 100))),
                        "Value": list(range(100)),
                    }
                ),
            )
        ]
        plan = f.starter_plan(data)
        tables, _ = f.execute_plan(plan, data)
        blob = f.build_excel(plan, tables)
        sheet = load_workbook(BytesIO(blob))["Data 01"]
        self.assertEqual(sheet["A5"].data_type, "s")
        self.assertEqual(sheet["B104"].value, 99)
        self.assertEqual(len(tables["input1"].frame), 100)
        original = f.digest_inputs("brief", data)
        data[0][1].loc[99, "Value"] = 1000
        self.assertNotEqual(original, f.digest_inputs("brief", data))
        precise = [("precise.csv", pd.DataFrame({"value": [0.1234567890123]}))]
        original = f.digest_inputs("brief", precise)
        precise[0][1].loc[0, "value"] = 0.1234567890124
        self.assertNotEqual(original, f.digest_inputs("brief", precise))

    def test_rejects_code_zero_division_unknown_columns_and_hidden_points(self):
        data, plan = scenario()
        for expression in [
            "__import__('os').system('echo bad')",
            "col('People').sum()",
            "2 ** 1000000",
            "col('Missing')",
            "col('People') / 0",
        ]:
            changed = deepcopy(plan)
            changed["tables"][0]["columns"][1]["expression"] = expression
            with (
                self.subTest(expression=expression),
                self.assertRaises((f.PlanError, ValueError)),
            ):
                f.execute_plan(changed, data)
        changed = deepcopy(plan)
        changed["charts"][2]["y_max"] = 5
        with self.assertRaisesRegex(f.PlanError, "hide observations"):
            f.execute_plan(changed, data)

    def test_invalid_join_and_histogram_coverage(self):
        data, plan = scenario()
        data[0][1].loc[1, "Year"] = 2001
        with self.assertRaises(ValueError):
            f.execute_plan(plan, data)
        data, plan = scenario()
        data[1][1].rename(columns={"2005/06": "2006/07"}, inplace=True)
        plan["tables"][1]["value_columns"][-1] = "2006/07"
        _, checks = f.execute_plan(plan, data)
        self.assertTrue(
            any(
                c["status"] == "warning" and "join coverage" in c["check"]
                for c in checks
            )
        )

    def test_actual_image_call_contract_crop_and_deck_evidence(self):
        picture = BytesIO()
        Image.new("RGB", (120, 100), "green").save(picture, "PNG")
        client = SimpleNamespace(
            images=SimpleNamespace(
                generate=Mock(
                    return_value=SimpleNamespace(
                        data=[
                            SimpleNamespace(
                                b64_json=base64.b64encode(picture.getvalue()).decode()
                            )
                        ]
                    )
                )
            )
        )
        request = {
            "id": "supply",
            "title": "Supply illustration",
            "prompt": "Illustration of a supply chain",
            "caption": "Illustrative context only.",
        }
        generated = f.generate_illustration(client, request)
        client.images.generate.assert_called_once()
        self.assertEqual(client.images.generate.call_args.kwargs["n"], 1)
        image = f.prepare_illustration(generated, 0.5, 0.5)
        self.assertEqual(image["original"], picture.getvalue())
        with Image.open(BytesIO(image["edited"])) as cropped:
            self.assertEqual(cropped.size, (120, 68))
        data, plan = scenario()
        plan["images"] = [request]
        tables, _ = f.execute_plan(plan, data)
        workbook = f.build_excel(plan, tables)
        with self.assertRaisesRegex(f.PlanError, "Generate illustration"):
            f.build_powerpoint(plan, tables, workbook, {})
        deck = f.build_powerpoint(plan, tables, workbook, {"supply": image})
        f.validate_files(plan, tables, workbook, deck)
        prs = Presentation(BytesIO(deck))
        text = "\n".join(
            s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame
        )
        self.assertIn("OpenAI imagegen", text)
        self.assertIn("Cropped in Study Buddy", text)
        self.assertNotIn(request["prompt"], text)
        self.assertEqual(len(prs.slides), 14)

    def test_planner_json_and_reference_validation(self):
        data, plan = scenario()
        client = SimpleNamespace(
            chat=SimpleNamespace(
                completions=SimpleNamespace(
                    create=Mock(
                        return_value=SimpleNamespace(
                            choices=[
                                SimpleNamespace(
                                    message=SimpleNamespace(content=json.dumps(plan))
                                )
                            ]
                        )
                    )
                )
            )
        )
        self.assertEqual(f.propose_plan(client, "test", "brief", data)["version"], 1)
        plan["sources"] = [
            {"title": "Provided evidence", "url": "https://example.test/evidence"}
        ]
        with self.assertRaises(f.PlanError):
            f.validate_sources(plan, "brief", [])
        f.validate_sources(
            plan, "brief", [("reference.md", "https://example.test/evidence")]
        )
        plan["images"] = [{"id": "same", "prompt": "a"}, {"id": "same", "prompt": "b"}]
        with self.assertRaises(f.PlanError):
            f.parse_plan(json.dumps(plan))


if __name__ == "__main__":
    unittest.main()
