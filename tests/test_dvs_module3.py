"""Module 3 requirements, original dataset controls and all guided handoffs."""
from pathlib import Path
import os,sys,tempfile,unittest,json,hashlib,zipfile
from types import SimpleNamespace
from unittest.mock import patch
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from streamlit.testing.v1 import AppTest
from openpyxl import load_workbook
import dvs_module3 as module
from dvs_curriculum import get_dvs_course_lessons,get_dvs_training_modules,get_dvs_task_guides,TEMPLATE_ROOT

class Module3Test(unittest.TestCase):
    def setUp(self):
        self.old=Path.cwd();self.temp=tempfile.TemporaryDirectory();os.chdir(self.temp.name)
    def tearDown(self):os.chdir(self.old);self.temp.cleanup()
    def run_app(self,app):
        app.run(timeout=90);self.assertFalse(app.exception,[e.message for e in app.exception]);return app
    def app(self,page='Learn & Practice'):
        app=AppTest.from_file(str(ROOT/'app.py'))
        for k,v in {'nav_section':'Learning','nav_page':page,'last_selected_course':'Data Visualisation','lp_selected_course':'FI1BBDV75 - Data Visualisation','tc_course_filter':'FI1BBDV75 - Data Visualisation'}.items():app.session_state[k]=v
        return self.run_app(app)
    def test_complete_coverage_and_source_controls(self):
        lessons=[l for l in get_dvs_course_lessons() if l['module_number']==3]
        self.assertEqual([l['lesson_number'] for l in lessons],['3.1','3.2','3.3','3.4'])
        self.assertEqual(len(module.GUIDES),18)
        self.assertEqual(len({g['uid'] for g in module.GUIDES}),18)
        self.assertEqual(len(get_dvs_task_guides('3.1')),13)
        self.assertEqual(len(get_dvs_task_guides('3.4')),3)
        for l in lessons:
            self.assertTrue((module.ROOT/l['dvs_illustrations'][0]['file']).is_file())
        for g in module.GUIDES:
            self.assertEqual((TEMPLATE_ROOT/f"DVS-{g['id']}-working-template.md").read_text(),module.get_template(g))
            self.assertIn(g['title'],g['markdown'])
        for source in module.WORKBOOKS:
            self.assertEqual(hashlib.sha256((module.DATA_ROOT/source['file']).read_bytes()).hexdigest(),source['sha256'])
        wb=load_workbook(module.DATA_ROOT/'C7M3LT_Data.xlsx',data_only=True,read_only=True)
        self.assertEqual(wb.sheetnames,['L2','L3','L4'])
        self.assertEqual(list(wb['L2'].values)[3][1:5],('Monday',20,12,47))
        self.assertEqual(wb['L3']['C3'].value,157);self.assertEqual(wb['L3']['D3'].value,23894)
        values=[wb['L4'].cell(r,3).value for r in range(4,20)]
        bounds=[wb['L4'].cell(r,4).value for r in range(4,8)]
        counts=[];previous=float('-inf')
        for upper in bounds:counts.append(sum(previous<v<=upper for v in values));previous=upper
        counts.append(sum(v>previous for v in values))
        self.assertEqual(counts,[5,6,1,4,0]);self.assertEqual(sum(counts),16)
        wb.close()
        wb=load_workbook(module.DATA_ROOT/'C7M3L5_Data.xlsx',data_only=True,read_only=True)
        for r in range(3,9):
            s=wb['Sheet1'];self.assertEqual(s.cell(r,7).value,s.cell(r,3).value*s.cell(r,4).value)
        wb.close()
        with zipfile.ZipFile(TEMPLATE_ROOT/'DVS Task 3.1 Best Practice Template.docx') as z:
            xml=z.read('word/document.xml').decode();styles=z.read('word/styles.xml').decode()
        for expected in ['three to four page','Why visualisation matters','How to identify and shape the data','Pitfalls and practical corrections','How to select and design a plot','Conclusion','References']:self.assertIn(expected,xml)
        self.assertNotIn('<w:pBdr',xml+styles)
    def test_four_tasks_and_fourteen_activities_transfer_independently(self):
        app=self.app()
        for guide in module.GUIDES:
            number=guide['lesson_number'];kind=('Task '+number) if guide['id'].endswith('-task') else 'Activity '+guide['id']
            next(s for s in app.selectbox if s.label==f'Choose a DVS Lesson {number} activity or task').set_value(guide['uid']);self.run_app(app)
            draft=f'{kind}: my actual check notes are incomplete. I have not reopened the output yet.'
            next(t for t in app.text_area if t.label==f'Your {kind} draft and evidence').set_value(draft);self.run_app(app)
            next(s for s in app.selectbox if s.label==f'Choose a {kind} step' and not s.key.startswith('course_exam_connector')).set_value(1);self.run_app(app)
            next(b for b in app.button if b.label==f'Prepare {kind} in Exam Resolver').click();self.run_app(app)
            prompt=next(t for t in app.text_area if t.label=='Paste the exam question, case, or task prompt').value
            self.assertIn('DVS '+kind+':',prompt);self.assertIn(draft,prompt)
            self.assertNotIn('Current selected step:',prompt)
            self.assertEqual([e.label for e in app.expander if e.label.startswith('Resolver task')],[f"Resolver task {guide['id']}: {guide['title']}"])
            self.assertEqual(next(s for s in app.selectbox if s.label==f'Choose a {kind} step' and s.key.startswith('course_exam_connector')).value,1)
            self.assertIn(f'Download {kind} working template (Markdown)',[d.label for d in app.get('download_button')])
            if guide['task_dataset']:self.assertIn(f"Download {guide['task_dataset']} (original source data)",[d.label for d in app.get('download_button')])
        self.assertNotIn('resolver_answer_history',app.session_state)
        self.assertEqual(len([uid for uid in app.session_state['dvs_persona_drafts'] if uid.startswith('dvs-3.')]),18)
    def test_english_normal_and_strict_selected_step_context(self):
        requests=[]
        def complete(**kwargs):
            requests.append(kwargs);return iter([SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content='English chart checklist review.'))])])
        fake=SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=complete)))
        env={'OPENAI_API_KEY':'offline-test','AI_PROVIDER':'openai'}
        env['_STUDY_BUDDY_INITIAL_ENV_KEYS']=json.dumps(list(os.environ)+list(env))
        with patch.dict(os.environ,env),patch('openai.OpenAI',return_value=fake):
            app=self.app()
            next(t for t in app.text_area if t.label=='Your Task 3.4 draft and evidence').set_value('I have not checked category-name removal yet.');self.run_app(app)
            next(b for b in app.button if b.label=='Prepare Task 3.4 in Exam Resolver').click();self.run_app(app)
            self.assertEqual(len(requests),0)
            next(s for s in app.selectbox if s.label=='Choose a Task 3.4 step' and s.key.startswith('course_exam_connector')).set_value(5);self.run_app(app)
            next(b for b in app.button if b.label=='💡 Generate answer with reasoning').click();self.run_app(app)
            next(c for c in app.checkbox if c.label=='📚 Strict curriculum mode (use ONLY the syllabus)').set_value(True);self.run_app(app)
            next(b for b in app.button if b.label=='🔁 Re-generate answer').click();self.run_app(app)
            self.assertEqual(len(requests),2)
            for r in requests:
                ctx=r['messages'][1]['content']
                for term in ['Current selected step: 6. Apply colour style and callout settings','I have not checked category-name removal yet','### DVS Task 3.4:','C4:C19','D4:D7','Style 6','Data Callout','Category Name','More 0','Unequal','0%']:
                    self.assertIn(term,ctx)
                self.assertNotIn('Current selected step: 1.',ctx)
                self.assertNotIn('### DVS Task 3.3:',ctx)
                self.assertIn('ANSWER LANGUAGE: Write in English',r['messages'][0]['content'])
            self.assertIn('STRICT CURRICULUM MODE',requests[1]['messages'][1]['content'])
            next(s for s in app.selectbox if s.label=='Choose a Task 3.4 step' and s.key.startswith('course_exam_connector')).set_value(0);self.run_app(app)
            self.assertTrue(any('earlier question or dataset' in i.value for i in app.info))
            self.assertEqual(len(requests),2)
    def test_training_module_and_unchanged_quiz_bank(self):
        data=get_dvs_training_modules()['DVS Module 3: Visualisation Techniques']
        self.assertEqual(len(data['lessons']),4);self.assertEqual(len(data['exercises']),18);self.assertEqual(len(data['quiz']),10)
        app=self.app('Training Center')
        next(s for s in app.selectbox if s.label=='🎯 Select Topic:').set_value('DVS Module 3: Visualisation Techniques');self.run_app(app)
        self.assertIn('Prepare Task 3.3 in Exam Resolver',[b.label for b in app.button])
        self.assertIn('Download Task 3.1 best-practice template (Word)',[d.label for d in app.get('download_button')])

if __name__=='__main__':unittest.main()
